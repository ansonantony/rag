import io, uuid
from pathlib import Path
from fastapi import UploadFile
from langchain_core.documents import Document as LCDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader
from qdrant_client import QdrantClient, models
from app.core.config import get_settings
from app.routes.query import HuggingFaceEmbeddings
from app.db.database import get_db
from app.db.models import Document
from sqlalchemy import select

ALLOWED = {".pdf": {"application/pdf"}, ".txt": {"text/plain"}, ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}}
def validate_upload(upload: UploadFile, content: bytes):
    ext = Path(upload.filename or "").suffix.lower()
    if ext not in ALLOWED: raise ValueError("Unsupported file type. Allowed types: PDF, TXT, DOCX.")
    if upload.content_type and upload.content_type not in ALLOWED[ext]: raise ValueError("File MIME type does not match an allowed document type.")
    if not content: raise ValueError("Empty files are not allowed.")
    if len(content) > get_settings().max_upload_size_mb * 1024 * 1024: raise OverflowError("File exceeds maximum upload size.")
    return ext
def extract_chunks(content: bytes, filename: str, ext: str):
    if ext == ".txt": docs = [LCDocument(page_content=content.decode("utf-8", errors="replace"), metadata={"page": 1})]
    elif ext == ".docx":
        from docx import Document as DocxDocument
        docs = [LCDocument(page_content="\n".join(p.text for p in DocxDocument(io.BytesIO(content)).paragraphs), metadata={"page": 1})]
    else:
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as file:
            file.write(content); path = file.name
        try: docs = PyPDFLoader(path).load()
        finally: Path(path).unlink(missing_ok=True)
    return RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200).split_documents(docs)
def ingest(document_id: uuid.UUID, user_id: uuid.UUID, filename: str, content: bytes, ext: str) -> int:
    settings = get_settings(); chunks = extract_chunks(content, filename, ext); embeddings = HuggingFaceEmbeddings()
    client = QdrantClient(url=settings.qdrant_url); vectors = embeddings.embed_documents([c.page_content for c in chunks])

    # Get the document to find its specific Qdrant collection
    from app.db.database import get_db
    db_gen = get_db()
    db = next(db_gen)
    try:
        document = db.scalar(select(Document).where(Document.id == document_id, Document.user_id == user_id))
        collection_name = document.qdrant_collection if document and document.qdrant_collection else settings.qdrant_collection_name

        # Create collection if it doesn't exist
        if not client.collection_exists(collection_name):
            client.create_collection(
                collection_name=collection_name,
                vectors_config=models.VectorParams(size=len(vectors[0]), distance=models.Distance.COSINE),
            )

        points = []
        for index, (chunk, vector) in enumerate(zip(chunks, vectors)):
            page = chunk.metadata.get("page_label", chunk.metadata.get("page", 0)); page = int(page) + 1 if isinstance(page, int) and page == chunk.metadata.get("page", 0) else page
            points.append(models.PointStruct(id=str(uuid.uuid4()), vector=vector, payload={"document_id": str(document_id), "user_id": str(user_id), "filename": filename, "chunk_id": f"chunk_{index}", "page_number": page, "source_text": chunk.page_content}))
        client.upsert(collection_name, points=points, wait=True)
    finally:
        db.close()
    return len(points)
